#!/usr/bin/env python3
"""
엑셀 시트 하나를 CSV 로 옮긴다.

    python scripts/xlsx-sheet-to-csv.py <엑셀> <시트이름> <나올CSV>

── 왜 필요한가 ─────────────────────────────────────────────

어휘를 확인하실 때 엑셀로 여시고 거기서 지우고 고치신다. 그런데 앱 쪽
도구(scripts/*.mjs)는 전부 node 로 돌고, node 는 엑셀을 못 읽는다.

그렇다고 도구마다 파이썬을 끼우면 어휘를 만질 때마다 두 가지를 갖춰야
한다. 그래서 **여기 한 곳에서만** 엑셀을 읽고, 나머지는 전부 CSV 로 다룬다.
국어 쪽(scripts/korean/xlsx-to-json.py)과 같은 생각이다.

맨 앞에 BOM 을 붙인다. 엑셀이 다시 열 때 한글이 안 깨진다.
"""

import csv
import io
import sys

import openpyxl


def cell(v) -> str:
    """
    칸 하나를 글자로.

    **엑셀이 참·거짓으로 바꿔 놓은 것을 되돌린다.** 어휘에 'true' 와 'false'
    가 낱말로 들어 있는데, 엑셀은 그 칸을 불리언으로 읽어 TRUE/FALSE 로
    바꾼다. 그대로 옮기면 파이썬이 'True'/'False' 로 적고, 그러면 앱의
    'true'/'false' 와 안 맞아 **멀쩡한 낱말을 뺀 것으로 잡는다.**
    실제로 그럴 뻔했다.
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    return str(v).strip()


def main() -> int:
    if len(sys.argv) < 4:
        print(__doc__)
        return 1

    src, sheet, out = sys.argv[1], sys.argv[2], sys.argv[3]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        print(f"  ✖ '{sheet}' 시트가 없습니다. 있는 시트 : {', '.join(wb.sheetnames)}")
        return 1

    ws = wb[sheet]
    rows = 0
    with io.open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        for row in ws.iter_rows(values_only=True):
            # 통째로 빈 줄은 버린다. 엑셀은 만졌던 자리를 빈 줄로 남긴다.
            cells = [cell(c) for c in row]
            if not any(cells):
                continue
            w.writerow(cells)
            rows += 1
    wb.close()

    print(f"  {src} [{sheet}] → {out}  ({rows}행)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
