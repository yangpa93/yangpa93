#!/usr/bin/env python3
"""넣은 낱말 묶음(data/batch-*.txt)을 **엑셀 한 장으로** 뽑는다.

    python scripts/batch-to-xlsx.py review/새낱말.xlsx data/batch-freq4-1.txt …

── 왜 필요한가 ─────────────────────────────────────────────

낱말을 새로 넣으면 회원님이 그것을 보고 "레벨이 이상하다" · "예문이 어색하다"
를 짚어 주셔야 한다. 그런데 넣는 형식(`낱말|품사|뜻|동의어|예문::해석|…`)은
기계가 읽자고 만든 것이라 사람이 훑기에 나쁘다. 마크다운으로도 뽑아 봤지만
**한 화면에 몇 개밖에 안 들어와** 224개를 보기가 어려웠다.

엑셀이면 걸러 보고(층·레벨), 정렬하고, 옆 칸에 바로 적을 수 있다.
어휘 표를 확인하실 때 늘 쓰시던 방식이기도 하다.

엑셀을 다루는 것은 파이썬 쪽에서만 한다 — scripts/xlsx-sheet-to-csv.py 와
같은 생각이다.
"""

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 화면에서 읽는 사람이 한국어라 한글이 곱게 나오는 글꼴을 쓴다.
FONT = "맑은 고딕"

TIER_NAME = {1: "기초 (초등 권장)", 2: "중급 (중학 권장)", 0: "고급 (고등)"}
LEVELS = [f"{g}-{s}" for g in ("m1", "m2", "m3", "h1", "h2", "h3") for s in (1, 2, 3, 4)]


def read_plan(path="src/data/plan.ts"):
    """배치표에서 `낱말 -> (레벨, 난이도 층)`."""
    out = {}
    level = None
    for line in Path(path).read_text(encoding="utf-8").split("\n"):
        head = re.match(r"^  '([a-z0-9-]+)': \[$", line.rstrip("\r"))
        if head:
            level = head.group(1)
            continue
        row = re.match(r'^    \["((?:[^"\\]|\\.)*)", (\d)', line)
        if row and level:
            out[row.group(1).replace('\\"', '"')] = (level, int(row.group(2)))
    return out


def read_batch(path):
    """넣은 묶음 한 장. 한 줄이 뜻 하나다(같은 낱말이 여러 줄일 수 있다)."""
    rows = []
    for line in Path(path).read_text(encoding="utf-8").split("\n"):
        line = line.rstrip("\r")
        if not line.strip() or line.startswith("#"):
            continue
        cells = line.split("|")
        word, pos, meaning, syn = cells[0], cells[1], cells[2], cells[3]
        examples = []
        for ex in cells[4:]:
            en, _, ko = ex.partition("::")
            examples.append((en, ko))
        rows.append((word, pos, meaning, syn, examples))
    return rows


def guide_sheet(wb):
    """무엇을 어떻게 봐 주셔야 하는지. **적는 칸이 있는 표에는 안내를 붙인다.**"""
    ws = wb.create_sheet("읽어 주세요")
    ws.sheet_view.showGridLines = False
    text = [
        ("곰탱이보카 — 새로 넣은 낱말 살펴보기", True),
        ("", False),
        ("뜻은 빈도4 시트에 적힌 것을 그대로 옮겼습니다.", False),
        ("난이도 층과 예문은 제가 정한 것이라, 그 둘을 봐 주셔야 합니다.", False),
        ("", False),
        ("① 「레벨」 칸이 눈에 걸리는 낱말", True),
        ("   중학교 1학년 첫 레벨(m1-1)에 어려운 낱말이 앉아 있지 않은지 봅니다.", False),
        ("   레벨 칸으로 걸러 보시면 한 레벨씩 훑을 수 있습니다.", False),
        ("", False),
        ("② 예문이 아이가 읽을 수 있는 문장인지", True),
        ("   문장 셋 중 하나라도 어색하면 그 줄만 알려 주시면 됩니다.", False),
        ("", False),
        ("③ 고치실 것은 맨 오른쪽 「여기 적어 주세요」 칸에", True),
        ("   그 칸에만 적어 주세요. 나머지 칸은 손대지 않으셔도 됩니다.", False),
        ("   보기 —  레벨이 너무 낮음. 중급으로", False),
        ("   보기 —  두 번째 예문이 어색함", False),
        ("", False),
        ("다 보시면 파일 그대로 돌려주시면 됩니다.", False),
    ]
    for i, (line, bold) in enumerate(text, start=2):
        c = ws.cell(row=i, column=2, value=line)
        c.font = Font(name=FONT, size=12 if i == 2 else 11, bold=bold)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 90
    return ws


# 검사기가 규칙으로 가릴 수 없어 **사람이 읽다 찾은** 것들.
# 「0개」 가 다 맞다는 뜻은 아니라는 표이기도 해서, 목록에 같이 싣는다.
EYE_FOUND = [
    ("distinction", "구별 명예", "쉼표가 빠졌습니다. `구별, 명예` 인 듯합니다"),
    ("trade-off", "이율 배반 · 절충 안", "`이율배반` · `절충안` 이 한 낱말입니다"),
    ("signature", "(고유 성을 잘 나타내는)", "`(고유성을`"),
    ("era", "(인생의) 한시기", "이쪽은 거꾸로 붙었습니다. `한 시기`"),
    ("ongoing", "진행중인", "이쪽도 거꾸로. `진행 중인`"),
]


def fix_sheet(wb, csv_path):
    """빈도4 표에서 아직 고칠 곳. **없으면 시트를 안 만든다.**

    엑셀에서 고치시는 김에 한 파일에서 다 보시는 편이 낫다. 다만 고치는 곳은
    원본(2027_영어_단어.xlsx)이지 이 표가 아니다 — 여기 것은 베껴 온 목록이다.
    """
    import json
    import subprocess

    try:
        got = subprocess.run(
            ["node", "scripts/check-meanings.mjs", csv_path, "--json"],
            capture_output=True, text=True, encoding="utf-8",
        )
        hits = json.loads(got.stdout)
    except Exception:
        return None
    if not hits:
        return None

    ws = wb.create_sheet(f"빈도4 고칠 곳 {len(hits) + len(EYE_FOUND)}")
    ws.append(["영단어", "무엇이 걸렸나", "걸린 자리", "지금 뜻 (표에 적힌 그대로)"])
    for h in hits:
        # 같은 까닭이 두 자리에서 걸리면 까닭은 한 번만 적는다.
        why = list(dict.fromkeys(w for w, _ in h["hits"]))
        ws.append([h["word"], " · ".join(why),
                   " / ".join(p for _, p in h["hits"]), h["meaning"]])
    for word, where, note in EYE_FOUND:
        ws.append([word, "눈으로 찾음 (기계가 못 잡는 모양)", where, note])

    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill("solid", fgColor="F8CBAD")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([18, 24, 40, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 1

    out_path, batches = sys.argv[1], sys.argv[2:]
    plan = read_plan()

    rows = []
    for b in batches:
        rows.extend(read_batch(b))

    # 레벨 순으로 놓는다. 한 레벨에 무엇이 들어갔는지가 이 표의 핵심이다.
    order = {lv: i for i, lv in enumerate(LEVELS)}
    rows.sort(key=lambda r: (order.get(plan.get(r[0], ("", 0))[0], 99), r[0]))

    wb = Workbook()
    wb.remove(wb.active)
    guide_sheet(wb)

    ws = wb.create_sheet(f"새 낱말 {len({r[0] for r in rows})}개")
    head = [
        "레벨", "난이도 층", "영단어", "품사", "뜻", "바꿔 쓸 수 있는 말",
        "예문 1", "해석 1", "예문 2", "해석 2", "예문 3", "해석 3",
        "여기 적어 주세요",
    ]
    ws.append(head)

    for word, pos, meaning, syn, examples in rows:
        level, tier = plan.get(word, ("?", -1))
        line = [level, TIER_NAME.get(tier, "?"), word, pos, meaning,
                ", ".join(s for s in syn.split(";") if s)]
        for en, ko in examples[:3]:
            line += [en, ko]
        line += [""] * (len(head) - len(line))
        ws.append(line)

    # ── 보기 좋게 ────────────────────────────────────────────
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    note_fill = PatternFill("solid", fgColor="FFF2CC")  # 적으실 칸은 노랗게
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "C2"  # 레벨과 층은 옆으로 밀려도 남는다
    ws.auto_filter.ref = ws.dimensions

    widths = [8, 16, 18, 7, 34, 22, 42, 34, 42, 34, 42, 34, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    last = len(head)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[last - 1].fill = note_fill

    fix = fix_sheet(wb, "review/빈도4.csv")

    wb.save(out_path)
    print(f"  {out_path}  ({len(rows)}줄 · 낱말 {len({r[0] for r in rows})}개"
          + (f" · 고칠 곳 {fix.max_row - 1}" if fix else "") + ")")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
