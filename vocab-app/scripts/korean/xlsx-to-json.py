#!/usr/bin/env python3
"""
국어 어휘 엑셀을 korean/source.json 으로 옮긴다.

엑셀을 그때그때 읽지 않고 JSON 을 한 번 만들어 커밋해 두는 이유:

  * 레벨 파일을 만드는 스크립트는 Node 로 돌아간다. 엑셀을 읽으려면
    파이썬과 openpyxl 이 필요한데, 레벨을 다시 뽑을 때마다 그 둘을
    갖춰야 하는 것은 번거롭다.
  * 엑셀은 원본 그대로 두고 싶다. 손으로 고친 흔적이 JSON 쪽에 남으면
    무엇이 원본이고 무엇이 수정인지 diff 로 보인다.

원본 엑셀은 둘이다.

  Korean_essential_voca.xlsx   개념어 · 고전 · 수능
  사자성어_고유어.xlsx          사자성어 · 고유어

**사자성어는 뒤엣것이 원본이다.** 앞엣것에도 사자성어 시트가 있지만, 뜻과
예문을 다시 손보고 115개를 더한 것이 뒤엣것이라 그쪽을 쓴다. 난이도(Level
1~6)도 거기 매겨져 있어서, 레벨을 나눌 때 difficulty.json 보다 먼저 본다.

CSAT_Korean_Essential_Vocabulary_800.xlsx 의 800개는 순번·어휘명까지
Korean_essential_voca.xlsx 안에 그대로 들어 있음을 확인했다.

    python3 scripts/korean/xlsx-to-json.py
"""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl 이 필요합니다:  pip install openpyxl')

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / 'korean' / 'Korean_essential_voca.xlsx'
IDIOM_XLSX = ROOT / 'korean' / '사자성어_고유어.xlsx'
OUT = ROOT / 'korean' / 'source.json'

# 시트마다 머리글 줄 위치와 열 순서가 제각각이라 여기에 적어 둔다.
# (시트 이름, 머리글이 있는 줄, 뽑아 올 열 번호)
SHEETS = {
    'concept': ('개념어_200개', 1, {'no': 0, 'field': 1, 'word': 2, 'meaning': 3, 'example': 4}),
    'classic': ('고전문학 필수 어휘 150 개', 0, {'no': 0, 'word': 1, 'meaning': 2, 'field': 3, 'example': 4}),
}

# 사자성어·고유어는 파일이 따로다. 예문이 두 칸이라 열 이름도 다르다.
IDIOM_SHEET = ('사자성어_300개_수정', 0,
               {'level': 0, 'no': 1, 'hanja': 2, 'word': 3, 'meaning': 4, 'sim': 5,
                'example': 6, 'example2': 7})
NATIVE_SHEET = ('국어1등급어휘력_고유어', 0,
                {'no': 0, 'word': 1, 'meaning': 2, 'example': 3})

CSAT_SHEETS = [
    '수능 국어 필수 어휘 (1~200)',
    '수능 국어 필수 어휘(201~400)',
    '수능 국어 필수 어휘 (401~600)',
    '수능 국어 필수 어휘 (601~800)',
]
CSAT_COLS = {'no': 0, 'word': 1, 'hanja': 2, 'field': 3, 'meaning': 4, 'example': 5}


def cell(v):
    return '' if v is None else str(v).strip()


def read(wb, sheet, header_row, cols):
    ws = wb[sheet]
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= header_row:
            continue
        if all(c is None or str(c).strip() == '' for c in row):
            continue
        rec = {k: cell(row[c]) for k, c in cols.items()}
        rec['no'] = int(float(rec['no'])) if rec['no'] else 0
        out.append(rec)
    return out


def clean(s):
    """
    예문 한 줄을 다듬는다.

    엑셀에는 큰따옴표로 감싼 것이 271개 있다. 그대로 쓰면 화면에서 따옴표가
    겹쳐 보인다. 그리고 마크다운 울타리(```)가 통째로 들어간 칸이 하나 있어서
    (양상군자), 한글이 한 자도 없으면 예문이 아닌 것으로 보고 버린다.

    **셀 안의 줄바꿈을 반드시 없앤다.** 엑셀 칸에서 Alt+Enter 로 줄을 나눈
    것이 그대로 넘어오면, 레벨 파일(.ts)에 따옴표가 열린 채 줄이 바뀌어
    **파일이 통째로 깨진다.** 실제로 사면초가 예문 하나 때문에 여섯 스위트가
    컴파일조차 못 했다.
    """
    s = ' '.join(str(s).split())          # 줄바꿈·연속 공백을 한 칸으로
    s = s.strip().strip('"').strip('“”').strip("'").strip()
    if not any('가' <= ch <= '힣' for ch in s):
        return ''
    return s


def read_idioms():
    """
    사자성어. **예문이 두 칸이라 목록으로 담는다.**

    두 칸이 똑같은 것이 114개 있다(같은 문장을 복사해 두셨다). 그런 것은
    하나로 줄인다 — 같은 문장을 두 번 보여 줄 까닭이 없다.

    담는 모양은 `[{"t": "문장"}]` 이다. **글자 목록이 아니라 객체 목록이라야
    한다** — build-levels 가 예문을 쓸 때 `e.t` 를 찾기 때문에, 글자를 그냥
    넣으면 렌더링에서 통째로 걸러져 예문이 하나도 없는 어휘가 된다.
    실제로 그렇게 만들었다가 416개가 빈 채로 나왔고 시험이 잡았다.
    """
    wb = openpyxl.load_workbook(IDIOM_XLSX, read_only=True, data_only=True)
    sheet, hr, cols = IDIOM_SHEET
    rows = read(wb, sheet, hr, cols)
    natives = read(wb, *NATIVE_SHEET[:2], NATIVE_SHEET[2])
    wb.close()

    out = []
    for r in rows:
        if not r['word']:
            continue
        seen = []
        for key in ('example', 'example2'):
            t = clean(r.pop(key, ''))
            if t and t not in seen:
                seen.append(t)
        r['examples'] = [{'t': t} for t in seen]
        r['level'] = int(float(r['level'])) if r['level'] else 0
        out.append(r)

    nat = []
    for r in natives:
        if not r['word']:
            continue
        # 고유어 뜻은 「1. … 2. …」 로 줄이 나뉜 것이 있다. 한 줄로 잇는다.
        r['meaning'] = ' '.join(r['meaning'].split())
        t = clean(r.pop('example', ''))
        r['examples'] = [{'t': t}] if t else []
        nat.append(r)

    return out, nat


def main():
    if not XLSX.exists():
        sys.exit(f'원본 엑셀이 없습니다: {XLSX}')
    if not IDIOM_XLSX.exists():
        sys.exit(f'사자성어·고유어 엑셀이 없습니다: {IDIOM_XLSX}')

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    data = {}
    for key, (sheet, hr, cols) in SHEETS.items():
        data[key] = read(wb, sheet, hr, cols)

    csat = []
    for sheet in CSAT_SHEETS:
        csat += read(wb, sheet, 2, CSAT_COLS)
    data['csat'] = csat
    wb.close()

    data['idiom'], data['native'] = read_idioms()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps(data, ensure_ascii=False, indent=1) + '\n',
        encoding='utf-8',
    )

    total = sum(len(v) for v in data.values())
    for k, v in data.items():
        print(f'  {k:8} {len(v):>4}')
    print(f'  {"합계":8} {total:>4}  →  {OUT.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
